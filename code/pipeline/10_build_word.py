"""Build an editable Word version of the revised paper from its LaTeX source."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"E:\Research\iv_surface")
PAPER = ROOT / "paper"
OUTPUT = ROOT / "output" / "word" / "iv_surface_revised.docx"
ASSETS = ROOT / "tmp" / "docx_assets"

BLUE = "1F4E78"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def read(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def set_font(run, name: str, size: float | None = None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = BLACK


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size="5"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, fld_end])
    set_font(run, "Times New Roman", 9)


def citation_order() -> tuple[list[str], dict[str, int]]:
    bbl = read(PAPER / "main.bbl")
    keys = re.findall(r"\\bibitem\{([^}]+)\}", bbl)
    return keys, {key: i + 1 for i, key in enumerate(keys)}


def reference_numbers() -> dict[str, str]:
    aux = read(PAPER / "main.aux")
    return {
        key: value
        for key, value in re.findall(
            r"\\newlabel\{([^}]+)\}\{\{([^}]*)\}", aux
        )
    }


def replace_commands(text: str, cite_map: dict[str, int], ref_map: dict[str, str]) -> str:
    def cite_repl(match):
        nums = [cite_map[k.strip()] for k in match.group(1).split(",") if k.strip() in cite_map]
        return "[" + ", ".join(str(n) for n in nums) + "]"

    text = re.sub(r"\\cite\{([^}]+)\}", cite_repl, text)
    text = re.sub(r"\\ref\{([^}]+)\}", lambda m: ref_map.get(m.group(1), "?"), text)
    text = re.sub(r"\\label\{[^}]+\}", "", text)
    text = text.replace("~", " ")
    text = text.replace("---", "-").replace("--", "-")
    text = text.replace("``", "“").replace("''", "”")
    text = text.replace(r"\'e", "é").replace(r"\'E", "É")
    text = text.replace(r"\%", "%").replace(r"\&", "&").replace(r"\_", "_")
    text = text.replace(r"\,", " ").replace(r"\;", " ").replace(r"\!", "")
    text = text.replace(r"\ ", " ")
    text = text.replace("{,}", ",")
    text = text.replace(r"\textbackslash", "\\")
    return text


GREEK = {
    r"\tau": "τ", r"\lambda": "λ", r"\rho": "ρ", r"\eta": "η",
    r"\gamma": "γ", r"\theta": "θ", r"\varphi": "φ",
    r"\phi": "φ", r"\sigma": "σ", r"\partial": "∂",
}


def math_text(text: str) -> str:
    text = text.strip().strip("$")
    text = text.replace(r"\left", "").replace(r"\right", "")
    text = text.replace(r"\geq", "≥").replace(r"\ge", "≥")
    text = text.replace(r"\leq", "≤").replace(r"\le", "≤")
    text = text.replace(r"\forall", "∀").replace(r"\times", "×")
    text = text.replace(r"\to", "→").replace(r"\approx", "≈")
    text = text.replace(r"\in", "∈").replace(r"\pm", "±")
    for command, symbol in GREEK.items():
        text = text.replace(command, symbol)
    for _ in range(4):
        text = re.sub(r"\\frac\{([^{}]+)\}\{([^{}]+)\}", r"(\1)/(\2)", text)
        text = re.sub(r"\\sqrt\{([^{}]+)\}", r"√(\1)", text)
        text = re.sub(r"\\(?:mathrm|text|operatorname|mathcal)\{([^{}]+)\}", r"\1", text)
    text = re.sub(r"\^\{([^{}]+)\}", r"^\1", text)
    text = re.sub(r"_\{([^{}]+)\}", r"_\1", text)
    text = text.replace(r"\quad", "  ").replace(r"\Big", "").replace(r"\big", "")
    text = text.replace(r"\overline", "mean")
    text = text.replace(r"\max", "max").replace(r"\sum", "Σ").replace(r"\ln", "ln")
    text = text.replace(r"\log", "log").replace(r"\sqrt", "√")
    text = text.replace(r"\text", "").replace(r"\mathrm", "")
    text = text.replace("\\", "")
    text = text.replace("{", "").replace("}", "")
    return re.sub(r"\s+", " ", text).strip()


TOKEN_RE = re.compile(
    r"(\$[^$]+\$|\\(?:emph|textbf|texttt)\{[^{}]*\})",
    re.DOTALL,
)


def clean_plain(text: str) -> str:
    text = re.sub(r"\\(?:footnotesize|centering|noindent)\b", "", text)
    text = re.sub(r"\\setlength\{[^}]+\}\{[^}]+\}", "", text)
    text = re.sub(r"\\(?:vspace|hspace)\*?\{[^}]+\}", "", text)
    text = re.sub(r"\\[A-Za-z]+\*?", "", text)
    text = text.replace("{", "").replace("}", "")
    return re.sub(r"\s+", " ", text).strip()


def clean_segment(text: str) -> str:
    """Clean LaTeX prose while retaining spaces around rich-text tokens."""
    if not text:
        return ""
    leading = " " if text[0].isspace() else ""
    trailing = " " if text[-1].isspace() else ""
    core = clean_plain(text)
    if not core:
        return " " if leading or trailing else ""
    return leading + core + trailing


def add_rich_text(paragraph, text: str, cite_map, ref_map, size=10.5):
    text = replace_commands(text, cite_map, ref_map)
    # Preserve inline mathematics inside simple emphasis commands. The outer
    # emphasis is secondary to keeping the mathematical content intact.
    for command in ("textbf", "emph"):
        text = re.sub(
            rf"\\{command}\{{([^{{}}]*\$[^$]+\$[^{{}}]*)\}}",
            r"\1",
            text,
        )
    pos = 0
    for match in TOKEN_RE.finditer(text):
        if match.start() > pos:
            plain = clean_segment(text[pos:match.start()])
            if plain:
                set_font(paragraph.add_run(plain), "Times New Roman", size)
        token = match.group(0)
        if token.startswith("$"):
            run = paragraph.add_run(math_text(token))
            set_font(run, "Cambria Math", size)
        else:
            cmd, inner = re.match(r"\\(emph|textbf|texttt)\{(.*)\}", token, re.DOTALL).groups()
            run = paragraph.add_run(clean_plain(inner))
            if cmd == "texttt":
                set_font(run, "Courier New", max(size - 0.5, 8.0))
            else:
                set_font(run, "Times New Roman", size, bold=cmd == "textbf", italic=cmd == "emph")
        pos = match.end()
    if pos < len(text):
        plain = clean_segment(text[pos:])
        if plain:
            set_font(paragraph.add_run(plain), "Times New Roman", size)


def add_body_paragraph(doc, text, cite_map, ref_map):
    text = re.sub(r"(?m)^%.*$", "", text).strip()
    if not text:
        return
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.22)
    add_rich_text(p, text.replace("\n", " "), cite_map, ref_map)


def add_equation(doc, latex: str, number: int):
    latex = re.sub(r"\\label\{[^}]+\}", "", latex)
    latex = re.sub(r"\s+", " ", latex).strip().rstrip(",.")
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.1)
    table.columns[1].width = Inches(0.45)
    set_table_borders(table, color="FFFFFF", size="0")
    equation = table.cell(0, 0).paragraphs[0]
    equation.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation.paragraph_format.space_before = Pt(3)
    equation.paragraph_format.space_after = Pt(3)
    set_font(equation.add_run(math_text(latex)), "Cambria Math", 11)
    marker = table.cell(0, 1).paragraphs[0]
    marker.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    marker.paragraph_format.space_before = Pt(3)
    marker.paragraph_format.space_after = Pt(3)
    set_font(marker.add_run(f"({number})"), "Times New Roman", 10)


def table_rows(block: str, cite_map, ref_map) -> list[list[str]]:
    match = re.search(r"(?s)\\begin\{tabular\}\{[^}]+\}(.*?)\\end\{tabular\}", block)
    if not match:
        return []
    body = re.sub(r"\\(?:toprule|midrule|bottomrule)", "", match.group(1))
    rows = []
    for row in re.split(r"\\\\", body):
        row = row.strip()
        if not row or "&" not in row:
            continue
        cells = []
        for cell in row.split("&"):
            cell = replace_commands(cell.strip(), cite_map, ref_map)
            cell = re.sub(r"\$([^$]+)\$", lambda m: math_text(m.group(1)), cell)
            cell = clean_plain(cell)
            cells.append(cell)
        rows.append(cells)
    return rows


def add_table(doc, block: str, number: int, cite_map, ref_map):
    caption = re.search(r"(?s)\\caption\{(.*?)\}\s*\\label", block)
    p = doc.add_paragraph()
    p.style = doc.styles["Caption"]
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(f"Table {number}. "), "Times New Roman", 9.5, bold=True)
    if caption:
        add_rich_text(p, caption.group(1), cite_map, ref_map, size=9.5)

    rows = table_rows(block, cite_map, ref_map)
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_header(table.rows[0])
    usable = 6.8
    first = min(2.35, usable * 0.34) if cols > 2 else usable * 0.55
    remaining = (usable - first) / (cols - 1) if cols > 1 else usable
    widths = [first] + [remaining] * (cols - 1) if cols > 1 else [usable]
    for i, values in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            value = values[j] if j < len(values) else ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            if i == 0:
                set_cell_shading(cell, BLUE)
                set_font(run, "Times New Roman", 8.5, bold=True)
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                if i % 2 == 0:
                    set_cell_shading(cell, PALE_BLUE)
                set_font(run, "Times New Roman", 8.5, bold=("textbf" in value))
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)


def add_figure(doc, block: str, number: int, cite_map, ref_map):
    source = re.search(r"\\includegraphics(?:\[[^]]*\])?\{([^}]+)\}", block)
    caption = re.search(r"(?s)\\caption\{(.*?)\}\s*\\label", block)
    if not source:
        return
    image_path = PAPER / source.group(1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(image_path), width=Inches(6.4))
    cap = doc.add_paragraph()
    cap.style = doc.styles["Caption"]
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(cap.add_run(f"Figure {number}. "), "Times New Roman", 9.5, bold=True)
    if caption:
        add_rich_text(cap, caption.group(1), cite_map, ref_map, size=9.5)


BLOCK_RE = re.compile(
    r"(?s)(\\section\*?\{[^}]+\}|\\subsection\{[^}]+\}|"
    r"\\begin\{equation\}.*?\\end\{equation\}|"
    r"\\begin\{table\}(?:\[[^]]*\])?.*?\\end\{table\}|"
    r"\\begin\{figure\}(?:\[[^]]*\])?.*?\\end\{figure\})"
)


def add_latex_body(doc, source: str, cite_map, ref_map):
    section_no = 0
    subsection_no = 0
    equation_no = 0
    table_no = 0
    figure_no = 0
    pos = 0

    def add_text(text):
        text = re.sub(r"\\label\{[^}]+\}", "", text)
        for paragraph in re.split(r"\n\s*\n", text):
            add_body_paragraph(doc, paragraph, cite_map, ref_map)

    for match in BLOCK_RE.finditer(source):
        add_text(source[pos:match.start()])
        block = match.group(0)
        if block.startswith(r"\section"):
            title = re.search(r"\\section\*?\{([^}]+)\}", block).group(1)
            starred = block.startswith(r"\section*")
            if not starred:
                section_no += 1
            subsection_no = 0
            p = doc.add_paragraph(style="Heading 1")
            p.paragraph_format.keep_with_next = True
            heading = title if starred else f"{section_no} {title}"
            add_rich_text(p, heading, cite_map, ref_map, size=13)
        elif block.startswith(r"\subsection"):
            title = re.search(r"\\subsection\{([^}]+)\}", block).group(1)
            subsection_no += 1
            p = doc.add_paragraph(style="Heading 2")
            p.paragraph_format.keep_with_next = True
            add_rich_text(p, f"{section_no}.{subsection_no} {title}", cite_map, ref_map, size=11)
        elif block.startswith(r"\begin{equation}"):
            equation_no += 1
            latex = re.search(r"(?s)\\begin\{equation\}(.*?)\\end\{equation\}", block).group(1)
            add_equation(doc, latex, equation_no)
        elif block.startswith(r"\begin{table}"):
            table_no += 1
            add_table(doc, block, table_no, cite_map, ref_map)
        elif block.startswith(r"\begin{figure}"):
            figure_no += 1
            add_figure(doc, block, figure_no, cite_map, ref_map)
        pos = match.end()
    add_text(source[pos:])


def bibliography_entries() -> list[str]:
    bbl = read(PAPER / "main.bbl")
    matches = list(re.finditer(r"\\bibitem\{[^}]+\}", bbl))
    entries = []
    for i, match in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else bbl.index(r"\end{thebibliography}")
        entries.append(bbl[match.end():end].strip())
    return entries


def configure_document(doc: Document):
    section = doc.sections[0]
    doc.settings.odd_and_even_pages_header_footer = False
    section.different_first_page_header_footer = False
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(5)

    title = doc.styles["Title"]
    title.font.name = "Times New Roman"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = BLACK

    for name, size, bold, italic in (
        ("Heading 1", 13, True, False),
        ("Heading 2", 11, True, True),
        ("Caption", 9.5, False, False),
    ):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = italic
        style.font.color.rgb = BLACK
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(13)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(5)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(9)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(3)
    doc.styles["Caption"].paragraph_format.space_before = Pt(5)
    doc.styles["Caption"].paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(header.add_run("Arbitrage-Constrained Neural Interpolation of the Implied Volatility Surface"),
             "Times New Roman", 8.5, italic=True)
    add_page_number(section.footer.paragraphs[0])


def main():
    ASSETS.mkdir(parents=True, exist_ok=True)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    cite_keys, cite_map = citation_order()
    ref_map = reference_numbers()
    main_tex = read(PAPER / "main.tex")

    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Arbitrage-Constrained Neural Interpolation of the Implied Volatility Surface"
    doc.core_properties.author = "Pramukh Jain"
    doc.core_properties.subject = "NIFTY 50 implied volatility surface interpolation"

    title_text = re.search(r"(?s)\\title\{(.*?)\}\s*\\author", main_tex).group(1)
    title_text = clean_plain(title_text.replace("\n", " "))
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    set_font(p.add_run(title_text), "Times New Roman", 20, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    set_font(p.add_run("Pramukh Jain"), "Times New Roman", 11, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    set_font(p.add_run("[Affiliation, city, country]"), "Times New Roman", 9.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run("pramukhjain.9966@gmail.com"), "Times New Roman", 9.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("This work received no external funding."), "Times New Roman", 9, italic=True)

    abstract = re.search(r"(?s)\\begin\{abstract\}(.*?)\\end\{abstract\}", main_tex).group(1)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    set_font(p.add_run("Abstract  "), "Times New Roman", 10.5, bold=True)
    add_rich_text(p, abstract.replace("\n", " "), cite_map, ref_map, size=10.5)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    keywords = re.search(r"(?s)\\begin\{keywords\}(.*?)\\end\{keywords\}", main_tex).group(1)
    p = doc.add_paragraph()
    set_font(p.add_run("Index Terms  "), "Times New Roman", 9.5, bold=True)
    add_rich_text(p, keywords.replace("\n", " "), cite_map, ref_map, size=9.5)

    intro_start = main_tex.index(r"\section{Introduction}")
    intro_end = main_tex.index(r"\input{sections/data}")
    body = main_tex[intro_start:intro_end]
    for name in ("data.tex", "method.tex", "design.tex", "results.tex", "discussion.tex", "conclusion.tex"):
        body += "\n\n" + read(PAPER / "sections" / name)
    post_start = main_tex.index(r"\section*{Data Availability Statement}")
    post_end = main_tex.index(r"\bibliographystyle")
    body += "\n\n" + main_tex[post_start:post_end]
    add_latex_body(doc, body, cite_map, ref_map)

    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = True
    set_font(p.add_run("References"),
             "Times New Roman", 13, bold=True)
    for i, entry in enumerate(bibliography_entries(), start=1):
        entry = re.sub(r"\\hskip.*?\\relax", " ", entry, flags=re.DOTALL)
        entry = re.sub(r"\\newblock", " ", entry)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(3)
        set_font(p.add_run(f"[{i}] "), "Times New Roman", 9)
        add_rich_text(p, entry.replace("\n", " "), cite_map, ref_map, size=9)

    p = doc.add_paragraph(style="Heading 1")
    set_font(p.add_run("Author Biography"), "Times New Roman", 13, bold=True)
    p = doc.add_paragraph()
    set_font(p.add_run("Pramukh Jain. "), "Times New Roman", 10, bold=True)
    set_font(p.add_run(
        "His research interests include computational finance, option pricing, "
        "shape-constrained machine learning, and empirical evaluation of financial models."
    ), "Times New Roman", 10)

    doc.save(OUTPUT)
    print(f"Created {OUTPUT} ({OUTPUT.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
